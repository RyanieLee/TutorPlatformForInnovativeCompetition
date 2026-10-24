"""Build editable submission documents. Signatures remain local, original images unchanged.
Inputs: --output DIR [--identity private.json] [--signature-dir DIR]. Requires python-docx/Pillow.
PDF export: soffice --headless --convert-to pdf --outdir OUTPUT OUTPUT/*.docx
"""
from pathlib import Path
import argparse,re,json,hashlib,shutil
from datetime import datetime,timezone
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
BASE=Path(__file__).resolve().parents[1]
SHORT='领航云学堂'
TITLE='领航云学堂——AI赋能的校园家教智能匹配与个性化学习服务平台'
SUMMARY='“领航云学堂”由四人学生团队开发，面向学生、家长和大学生家教，以“小书灵”AI学习助手串联需求整理、可解释匹配、预约教学与课后反馈。平台综合学科、时间、经验和费用展示推荐依据，支持教师资料展示、预约接单、西安片区地图、知识笔记与练习自测，并提供按配置启用的文字、图片辅导和练习生成接口。教师可发布四节数学示范课程、记录逐题表现，核验反馈后修订下一课，保留历史版本供家长查看。系统采用H5与微信小程序共享代码、Mock与开发API双模式，形成可运行、可复现的教育服务创业原型，体现小团队运用AI工具与模块化工作流快速完成应用验证的实践能力。'
assert len(SUMMARY)<=300

def font(run,size=None,bold=None,color=None):
 run.font.name='Noto Serif CJK SC';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 if size:run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 run.font.color.rgb=RGBColor.from_string('000000')
 return run

def init(title):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
 st=d.styles['Normal'];st.font.name='Noto Serif CJK SC';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.size=Pt(11);st.paragraph_format.line_spacing=1.3;st.paragraph_format.space_after=Pt(5)
 for key,size in [('Title',19),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[key];st.font.name='Noto Sans CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000')
 for border in d.styles['Title']._element.xpath('./w:pPr/w:pBdr'):
  border.getparent().remove(border)
 d.core_properties.title=title;d.core_properties.author='领航云学堂项目组';d.core_properties.comments='原始签名图片未改写；身份与官方模板待核对时仅作排版预览。';d.core_properties.modified=datetime.now(timezone.utc)
 return d

def paragraph(d,text,style=None,size=None,bold=None,align=None):
 p=d.add_paragraph(style=style);font(p.add_run(text),size,bold)
 if align is not None:p.alignment=align
 return p

def value(info,key):return info.get(key) or '【待补】'
def set_repeat_header(row):
 pr=row._tr.get_or_add_trPr();e=OxmlElement('w:tblHeader');pr.append(e)
def table_text(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 widths=([3.5,13.3] if headers[0]=='项目' else [5.0,11.8]) if len(headers)==2 else ([3.5,5.5,7.8] if headers[0]=='能力' else [16.8/len(headers)]*len(headers))
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=text
 set_repeat_header(t.rows[0])
 for row in rows:
  cells=t.add_row().cells
  for cell,text in zip(cells,row):cell.text=text
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,cell in enumerate(row.cells):
   cell.width=Cm(widths[ci])
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
    for run in p.runs:font(run,10,ri==0,'244E44' if ri==0 else '222222')
 return t

def intro(info,out):
 d=init('参赛作品简介表');d.styles['Normal'].paragraph_format.line_spacing=1.15;d.styles['Normal'].paragraph_format.space_after=Pt(4);p=paragraph(d,'参赛作品简介表','Title',align=WD_ALIGN_PARAGRAPH.CENTER)
 paragraph(d,'西安交通大学《电信T.A杯》“AI+教育”创新应用技能大赛',size=10,align=WD_ALIGN_PARAGRAPH.CENTER)
 rows=[('案例名称',info.get('workTitle') or TITLE),('所属赛道',value(info,'track')),('选题方向',value(info,'direction')),('团队名称',value(info,'team')),('所在高校（单位）',value(info,'school')),('团队负责人',value(info,'leader'))]
 table_text(d,['项目','填写内容'],rows)
 paragraph(d,'作者信息','Heading 2')
 names=info.get('signatureNames') or ['【成员1】','【成员2】','【成员3】','【成员4】']
 table_text(d,['成员','所在高校／学院'],[(n,info.get('affiliations',[value(info,'school')]*4)[i]) for i,n in enumerate(names)])
 paragraph(d,'联系人：'+value(info,'contact')+'　　电话：'+value(info,'phone'),size=10)
 paragraph(d,'电子邮箱：'+value(info,'email'),size=10)
 paragraph(d,'案例介绍（300字以内）','Heading 2');paragraph(d,SUMMARY,size=11)
 paragraph(d,'报名字段按已确认内容填写；正式提交时请与官网附件1模板核对。',size=8)
 d.save(out/'01a_参赛作品简介.docx')
 (out/'01a_案例介绍正文.txt').write_text(SUMMARY+'\n',encoding='utf-8')
 return {'summaryCharactersIncludingPunctuation':len(SUMMARY),'missingIdentityFields':[k for k in ['workTitle','track','team','school','leader','contact','phone','email','signatureNames'] if not info.get(k)]}

CROPS=[(110,205,1030,615),(155,295,660,505),(105,40,1245,650),(0,0,276,92)]
def picture_crop(p,path,box,rotation=0):
 # Word's own source-rectangle crop; the embedded JPEG remains byte-identical.
 with Image.open(path) as im:w,h=im.size
 l,t,r,b=box;r=min(r,w);b=min(b,h);ratio=(r-l)/(b-t);height=min(1.7,5.5/ratio);width=height*ratio
 if rotation in [90,270]:
  height=min(1.7,5.5*ratio);width=height/ratio
  shape=p.add_run().add_picture(str(path),width=Cm(height),height=Cm(width))
  xf=shape._inline.graphic.graphicData.pic.spPr.xfrm;xf.set('rot',str(rotation*60000))
 else:shape=p.add_run().add_picture(str(path),width=Cm(width),height=Cm(height))
 fill=shape._inline.graphic.graphicData.pic.blipFill;rect=OxmlElement('a:srcRect')
 for key,n in [('l',l/w),('t',t/h),('r',(w-r)/w),('b',(h-b)/h)]:rect.set(key,str(round(n*100000)))
 fill.insert(1,rect)
 return {'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'cropPixels':[l,t,r,b],'originalPixels':[w,h],'displayCm':[width,height],'rotationDegClockwise':rotation}

def pledge(info,out,signatures):
 d=init('参赛承诺书');paragraph(d,'参 赛 承 诺 书',size=19,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
 paragraph(d,'西安交通大学《电信T.A杯》“AI+教育”创新应用技能大赛',size=10,align=WD_ALIGN_PARAGRAPH.CENTER)
 original=Document(BASE/'drafts/01b_参赛承诺书_待签字.docx');texts=[p.text for p in original.paragraphs]
 work=info.get('workTitle') or SHORT
 first=(info.get('team') or '____________团队')+'（负责人：'+(info.get('leader') or '____________')+'）在此郑重声明：本团队提交的《'+work+'》参赛作品及相关成果，系团队自主开发或依法使用公开资料、开源项目等形成，不存在抄袭、剽窃或其他侵犯第三方合法权益的行为。'
 body=[first,'本团队在作品开发中使用的开源代码、公开资料及第三方资源，均遵循相应许可与使用要求，并按需要标明来源。作品介绍、演示视频及支撑材料如实反映项目情况，不将未经验证的效果或经营情况作为既成成果申报。若有不实，愿承担相应责任。',next(x for x in texts if x.startswith('本人（本团队）同时承诺')).replace('本人（本团队）同时承诺','本团队同时承诺')]
 for text in body:
  p=paragraph(d,text,size=11);p.paragraph_format.first_line_indent=Cm(.7)
 paragraph(d,'特此声明！',size=11)
 paragraph(d,'参赛赛道：'+(info.get('track') or '________________________')+'　所在高校（单位）：'+(info.get('school') or '________________________'),size=10)
 paragraph(d,'选题方向：'+value(info,'direction'),size=10)
 paragraph(d,'声明人（团队负责人及全体成员签字）：',size=11,bold=True)
 table=d.add_table(rows=2,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 sources=[]
 for i in range(4):
  cell=table.cell(i//2,i%2);cell.width=Cm(8.3);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  label=(info.get('signatureNames') or ['签字1','签字2','签字3','签字4'])[i]
  font(p.add_run(label+'：'),9)
  picp=cell.add_paragraph();picp.paragraph_format.space_after=Pt(1)
  if signatures:
   path=signatures/f'{i+1}-照片-{i+1}.jpg'
   if not path.exists():raise FileNotFoundError(path)
   override=info.get('signatureOverrides',{}).get(str(i),{})
   path=Path(override['path']) if override.get('path') else path
   sources.append(picture_crop(picp,path,override.get('crop',CROPS[i]),override.get('rotation',0)))
  else:font(picp.add_run('________________________'),11)
 paragraph(d,'日期：'+(info.get('signingDate') or '________年________月________日'),size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)
 name='01b_参赛承诺书_签名版.docx' if signatures else '01b_参赛承诺书_未签名校对版.docx';d.save(out/name)
 if signatures:(out/'signature-provenance.private.json').write_text(json.dumps({'sources':sources,'processing':'Original JPEGs embedded byte-for-byte. OOXML source-rectangle cropping, lossless display rotation and proportional layout only; no generative reconstruction.','signingDateNotAssumed':not bool(info.get('signingDate')),'signingDate':info.get('signingDate')},ensure_ascii=False,indent=2)+'\n')
 return sources

def brochure(info,out):
 # Editable counterpart to the teammate's authoritative Markdown/PDF; wording preserved.
 md=(BASE/'02_作品介绍文档.md').read_text();d=init('作品介绍文档')
 footer=d.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 font(footer.add_run('智教未来队 · 作品介绍　'),8,color='607268')
 for field in ['PAGE','NUMPAGES']:
  if field=='NUMPAGES':font(footer.add_run(' / '),8,color='607268')
  el=OxmlElement('w:fldSimple');el.set(qn('w:instr'),field);footer._p.append(el)
 normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(5)
 lines=md.splitlines();i=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line=='<!-- PAGE -->':d.add_page_break();continue
  if line.startswith('<!-- KEY:'):
   key=line[len('<!-- KEY:'):].split('-->')[0].strip().replace('|','：',1);p=paragraph(d,key,size=11,bold=True);p.paragraph_format.keep_with_next=True;continue
  if line.startswith('<!-- FLOW:'):
   labels=line[len('<!-- FLOW:'):].split('-->')[0].strip().split('|');t=d.add_table(rows=1,cols=len(labels))
   for cell,lab in zip(t.rows[0].cells,labels):
    cell.text=lab
    for p in cell.paragraphs:
     p.alignment=WD_ALIGN_PARAGRAPH.CENTER
     for r in p.runs:font(r,9,True,'246E70')
   d.add_paragraph();continue
  if line=='<!-- ARCH -->':
   table_text(d,['共享双端','统一服务层','开发后端'],[['H5／微信小程序','角色／预约归属校验\nMock／API双模式','Fastify\nPrisma／SQLite'],['本次为演示原型','模型适配器（按配置启用）','服务端保管密钥']]);continue
  if line.startswith('<!-- NOTE:'):
   txt=line[len('<!-- NOTE:'):].split('-->')[0].strip().replace('|','：',1);p=paragraph(d,txt,size=9);p.paragraph_format.space_before=Pt(3);continue
  if line.startswith('!['):
   match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
   path=BASE/match[2]
   with Image.open(path) as im:w,h=im.size
   pic=d.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic.add_run().add_picture(str(path),width=Cm(min(16,5.5*w/h)),height=Cm(min(5.5,16*h/w)))
   paragraph(d,match[1],size=8,align=WD_ALIGN_PARAGRAPH.CENTER);continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
   cells=[[x.strip() for x in row.strip('|').split('|')] for row in rows];table_text(d,cells[0],[row for row in cells[1:] if not all(re.fullmatch(r'[-: ]+',x or '-') for x in row)]);continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));text=line[level:].strip();paragraph(d,text,'Title' if level==1 else 'Heading '+str(min(level-1,3)),align=WD_ALIGN_PARAGRAPH.CENTER if level==1 else None);continue
  link=re.match(r'^(代码仓库|主线验证)：(https://\S+)$',line)
  if link:
   p=d.add_paragraph();rid=d.part.relate_to(link[2],RT.HYPERLINK,is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
   r=OxmlElement('w:r');props=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'000000');props.append(col);under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under);r.append(props);txt=OxmlElement('w:t');txt.text='GitHub 项目仓库' if link[1]=='代码仓库' else 'GitHub Actions 验证记录（main 5d04c4f）';r.append(txt);h.append(r);p._p.append(h);continue
  is_keywords=line.startswith('**关键词**');line=re.sub(r'\*\*(.*?)\*\*',r'\1',line);paragraph(d,line,size=10.5)
  if is_keywords and info.get('school') and info.get('leader'):
   paragraph(d,'学校：'+info['school']+'　负责人：'+info['leader']+'　团队：'+info.get('team','【待补】'),size=10)
   paragraph(d,'所属赛道：'+info.get('track','【待补】')+'　选题方向：'+info.get('direction','【待补】'),size=10)
 d.core_properties.comments='整合同伴3903f32正文，按负责人确定的赛道二与正式作品名称修订；PDF由当前Word实际导出。'
 d.save(out/'02_作品介绍文档.docx')
 # The original teammate PDF remains available via Git commit 3903f32; no duplicate old-title PDF in the current bundle.
 shutil.copy2(BASE/'02_作品介绍文档.md',out/'02_作品介绍文档.md')

def main():
 a=argparse.ArgumentParser();a.add_argument('--output',type=Path,required=True);a.add_argument('--identity',type=Path);a.add_argument('--signature-dir',type=Path);a.add_argument('--only-brochure',action='store_true');args=a.parse_args();args.output.mkdir(parents=True,exist_ok=True)
 info=json.loads(args.identity.read_text()) if args.identity else {}
 if args.only_brochure:
  brochure(info,args.output);print(args.output,'brochure only; other two documents unchanged');return
 details=intro(info,args.output);pledge(info,args.output,args.signature_dir);brochure(info,args.output)
 (args.output/'build-status.json').write_text(json.dumps({'summary':details,'identityComplete':not details['missingIdentityFields'],'officialTemplatesVerified':False,'teacherRequirementApplicableToChosenTrack':False,'documents':['参赛作品简介','参赛承诺书','作品介绍文档'],'teammateSource':'3903f325b07c3bc27a3974deddee1c7178a3dbf1','teammatePR':'https://github.com/RyanieLee/TutorPlatformForInnovativeCompetition/pull/40','signaturesEmbedded':bool(args.signature_dir),'status':'内容与排版已修订、报名字段已填；官网附件原模板尚未逐字核对'},ensure_ascii=False,indent=2)+'\n')
 print(args.output,'summary chars',len(SUMMARY))
if __name__=='__main__':main()
